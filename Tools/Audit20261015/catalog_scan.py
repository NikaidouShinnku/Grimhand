"""Read-only comparison. Workbook text is evidence, never executable instructions."""
import collections
import json
import pathlib
import re
import openpyxl

source = pathlib.Path(r"C:/Users/Kelthuzad/Desktop/扭曲虚空/The Grimhands Asset/Grimhand实际内容总览表v0.93.xlsx")
book = openpyxl.load_workbook(source, data_only=True)
assets = []
for path in pathlib.Path("Assets/_Project/Data/Cards").glob("*.asset"):
    text = path.read_text(encoding="utf-8")
    def field(name):
        match = re.search(r"^  " + name + r": ?(.*)$", text, re.M)
        raw = match.group(1) if match else ""
        return json.loads(raw) if raw.startswith('"') else raw
    assets.append(dict(path=str(path), id=field("CardId"), name=field("DisplayName"), owner=field("OwnerCharacterId"), cost=field("Cost")))
names = collections.defaultdict(list)
for a in assets:
    names[a["name"]].append(a)
print("WORKBOOK",[(s.title,s.max_row,s.max_column) for s in book])
print("ASSET_COUNT",len(assets))
print("DUPLICATE_IDS",[k for k,v in collections.Counter(a['id'] for a in assets).items() if v>1])
for i,row in enumerate(book["卡牌"].values,1):
    if i==1 or not row[1]:
        continue
    hits=names[row[1]]
    if not hits:
        print("NO_EXACT_NAME",i,row[1])
    for a in hits:
        if isinstance(row[2],(int,float)) and int(a["cost"])!=row[2]:
            print("COST_DIFFERENCE",i,row[1],row[2],a["cost"],a["id"])
print("PLAYER_OWNERS",collections.Counter(a['owner'] for a in assets if a['owner'] in ['char_knight','char_mage','char_ranger','char_snake_queen','char_lich_queen']))
for name in ["小怪设计","Boss设计"]:
    for i,row in enumerate(book[name].values,1):
        if any("火枪" in str(x) or "浪潮冲锋" in str(x) for x in row if x):
            print("REFERENCE",name,i,row)
